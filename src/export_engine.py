import json
import smtplib
from datetime import datetime
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak,
    Image as RLImage
)
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT


def _safe_cell_value(value):
    """Prevent CSV-style formula execution when a workbook is opened."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def export_to_pdf(report: dict, output_path: Path = None) -> bytes:
    """Generate a professional PDF report with summary and exception details."""
    
    buffer = BytesIO() if output_path is None else None
    pdf_path = output_path or buffer
    
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    
    styles = getSampleStyleSheet()
    story = []
    
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1F3A93'),
        spaceAfter=6,
        fontName='Helvetica-Bold',
    )
    story.append(Paragraph('Ledger', title_style))
    
    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#6B7280'),
        spaceAfter=12,
    )
    story.append(Paragraph('AP Reconciliation Report', subtitle_style))
    
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    story.append(Paragraph(f'Generated: {timestamp}', styles['Normal']))
    story.append(Spacer(1, 0.2 * inch))
    
    summary_title = ParagraphStyle(
        'SectionTitle',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#1F3A93'),
        spaceAfter=10,
        fontName='Helvetica-Bold',
    )
    story.append(Paragraph('Executive Summary', summary_title))
    
    s = report['summary']
    summary_data = [
        ['Metric', 'Value'],
        ['Total Lines Processed', str(s['total_lines_processed'])],
        ['Total Exceptions Found', str(s['total_exceptions'])],
        ['Total Amount at Risk', f"${s['total_amount_at_risk']:,.2f}"],
        ['Touchless Rate', f"{s['touchless_rate_pct']}%"],
        ['Needing Review', str(s['needs_review'])],
        ['Auto-Resolved', str(s['auto_resolved'])],
    ]
    
    summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F3A93')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 0.3 * inch))
    
    story.append(Paragraph('Exception Details', summary_title))
    exceptions = report.get('exceptions', [])
    
    if exceptions:
        by_severity = {}
        for exc in exceptions:
            sev = exc['severity']
            if sev not in by_severity:
                by_severity[sev] = []
            by_severity[sev].append(exc)
        
        for severity in ['Critical', 'High', 'Medium', 'Low']:
            if severity not in by_severity:
                continue
            
            story.append(Spacer(1, 0.15 * inch))
            severity_color = {
                'Critical': '#B71C1C',
                'High': '#D32F2F',
                'Medium': '#FF9F1C',
                'Low': '#00A86B',
            }.get(severity, '#6B7280')
            
            sev_style = ParagraphStyle(
                f'{severity}Heading',
                parent=styles['Heading3'],
                fontSize=11,
                textColor=colors.HexColor(severity_color),
                spaceAfter=8,
                fontName='Helvetica-Bold',
            )
            story.append(Paragraph(f'{severity.upper()} Severity ({len(by_severity[severity])} items)', sev_style))
            
            for exc in by_severity[severity][:5]:
                exc_style = ParagraphStyle(
                    'ExcDetail',
                    parent=styles['Normal'],
                    fontSize=9,
                    leftIndent=0.2 * inch,
                    spaceAfter=6,
                )
                vendor = exc.get('vendor_name', 'Unknown')
                exc_type = exc.get('exception_type', '').replace('_', ' ')
                amount = f"${exc.get('amount_at_risk', 0):,.2f}"
                story.append(Paragraph(
                    f'<b>{vendor}</b> • {exc_type} • {amount}',
                    exc_style
                ))
                
                ai_summary = exc.get('ai_summary', '')
                if ai_summary:
                    story.append(Paragraph(ai_summary[:200] + '...', styles['Normal']))
                story.append(Spacer(1, 0.08 * inch))
        
        if len(exceptions) > 5:
            story.append(Paragraph(f'... and {len(exceptions) - 5} more exceptions', styles['Italic']))
    
    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph(
        'For complete exception details and policy citations, please review the interactive dashboard or Excel export.',
        styles['Normal']
    ))
    
    doc.build(story)
    
    if output_path is None:
        return buffer.getvalue()


def export_to_excel(report: dict, output_path: Path = None) -> bytes:
    """Generate a detailed Excel spreadsheet with summary and all exceptions."""
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    ws_summary = wb.create_sheet('Summary', 0)
    ws_summary['A1'] = 'Ledger — AP Reconciliation Report'
    ws_summary['A1'].font = Font(size=16, bold=True, color='FFFFFF')
    ws_summary['A1'].fill = PatternFill(start_color='1F3A93', end_color='1F3A93', fill_type='solid')
    ws_summary.merge_cells('A1:D1')
    ws_summary['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws_summary.row_dimensions[1].height = 25
    
    ws_summary['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_summary['A2'].font = Font(size=10, italic=True, color='6B7280')
    
    ws_summary['A4'] = 'Metric'
    ws_summary['B4'] = 'Value'
    for cell in ['A4', 'B4']:
        ws_summary[cell].font = Font(bold=True, color='FFFFFF')
        ws_summary[cell].fill = PatternFill(start_color='2E5BCC', end_color='2E5BCC', fill_type='solid')
        ws_summary[cell].alignment = Alignment(horizontal='left')
    
    s = report['summary']
    summary_rows = [
        ('Total Lines Processed', s['total_lines_processed']),
        ('Total Exceptions Found', s['total_exceptions']),
        ('Total Amount at Risk', s['total_amount_at_risk']),
        ('Touchless Rate (%)', s['touchless_rate_pct']),
        ('Exceptions Needing Review', s['needs_review']),
        ('Auto-Resolved', s['auto_resolved']),
        ('Processing Time (sec)', report.get('run_metadata', {}).get('processing_seconds', 'N/A')),
    ]
    
    for idx, (metric, value) in enumerate(summary_rows, start=5):
        ws_summary[f'A{idx}'] = metric
        ws_summary[f'B{idx}'] = value
        if idx % 2 == 0:
            ws_summary[f'A{idx}'].fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
            ws_summary[f'B{idx}'].fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
    
    ws_summary.column_dimensions['A'].width = 30
    ws_summary.column_dimensions['B'].width = 20
    
    ws_exc = wb.create_sheet('Exceptions', 1)
    
    headers = [
        'Exception ID', 'Vendor', 'Type', 'Severity', 'Status', 'Amount at Risk',
        'PO ID', 'Invoice ID', 'Payment ID', 'AI Summary', 'Recommended Action',
        'Policy Clause', 'Clause Text'
    ]
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws_exc.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1F3A93', end_color='1F3A93', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    severity_colors = {
        'Critical': 'FFCDD2',
        'High': 'FFEBEE',
        'Medium': 'FFF4E1',
        'Low': 'D4F8E8',
    }
    
    for row_idx, exc in enumerate(report.get('exceptions', []), start=2):
        ws_exc.cell(row=row_idx, column=1).value = _safe_cell_value(exc.get('exception_id', ''))
        ws_exc.cell(row=row_idx, column=2).value = _safe_cell_value(exc.get('vendor_name', ''))
        ws_exc.cell(row=row_idx, column=3).value = _safe_cell_value(exc.get('exception_type', '').replace('_', ' '))
        ws_exc.cell(row=row_idx, column=4).value = _safe_cell_value(exc.get('severity', ''))
        ws_exc.cell(row=row_idx, column=5).value = _safe_cell_value(exc.get('status', ''))
        ws_exc.cell(row=row_idx, column=6).value = exc.get('amount_at_risk', 0)
        ws_exc.cell(row=row_idx, column=7).value = _safe_cell_value(exc.get('po_id', ''))
        ws_exc.cell(row=row_idx, column=8).value = _safe_cell_value(exc.get('invoice_id', ''))
        ws_exc.cell(row=row_idx, column=9).value = _safe_cell_value(exc.get('payment_id', ''))
        ws_exc.cell(row=row_idx, column=10).value = _safe_cell_value(exc.get('ai_summary', ''))
        ws_exc.cell(row=row_idx, column=11).value = _safe_cell_value(exc.get('recommended_action', ''))
        ws_exc.cell(row=row_idx, column=12).value = _safe_cell_value(f"Clause {exc.get('cited_clause_id', '')}")
        ws_exc.cell(row=row_idx, column=13).value = _safe_cell_value(exc.get('cited_clause_text', ''))
        
        severity = exc.get('severity', '')
        if severity in severity_colors:
            for col_idx in range(1, len(headers) + 1):
                ws_exc.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    start_color=severity_colors[severity],
                    end_color=severity_colors[severity],
                    fill_type='solid'
                )
        
        ws_exc.cell(row=row_idx, column=6).number_format = '$#,##0.00'
        
        for col_idx in [3, 10, 11, 13]:
            ws_exc.cell(row=row_idx, column=col_idx).alignment = Alignment(wrap_text=True)
    
    ws_exc.column_dimensions['A'].width = 14
    ws_exc.column_dimensions['B'].width = 18
    ws_exc.column_dimensions['C'].width = 18
    ws_exc.column_dimensions['D'].width = 12
    ws_exc.column_dimensions['E'].width = 14
    ws_exc.column_dimensions['F'].width = 15
    ws_exc.column_dimensions['G'].width = 12
    ws_exc.column_dimensions['H'].width = 12
    ws_exc.column_dimensions['I'].width = 12
    ws_exc.column_dimensions['J'].width = 25
    ws_exc.column_dimensions['K'].width = 25
    ws_exc.column_dimensions['L'].width = 14
    ws_exc.column_dimensions['M'].width = 30
    
    ws_exc.freeze_panes = 'A2'
    
    if output_path:
        wb.save(output_path)
        return None
    else:
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def export_to_json(report: dict, output_path: Path = None) -> bytes:
    """Export report as JSON for integration with other systems."""
    
    json_str = json.dumps(report, indent=2, default=str)
    
    if output_path:
        output_path.write_text(json_str)
        return None
    else:
        return json_str.encode('utf-8')


def send_email_report(
    report: dict,
    recipient_email: str,
    sender_email: str = None,
    sender_password: str = None,
    smtp_server: str = 'smtp.gmail.com',
    smtp_port: int = 587,
) -> dict:
    """
    Send report via email with attachments.
    
    Args:
        report: Reconciliation report dict
        recipient_email: Email address to send to
        sender_email: Gmail address (default: from .env SENDER_EMAIL)
        sender_password: Gmail app password (default: from .env SENDER_PASSWORD)
        smtp_server: SMTP server address
        smtp_port: SMTP port
        
    Returns:
        {'success': bool, 'message': str}
    """
    
    if not sender_email or not sender_password:
        from dotenv import load_dotenv
        import os
        load_dotenv()
        sender_email = sender_email or os.getenv('SENDER_EMAIL')
        sender_password = sender_password or os.getenv('SENDER_PASSWORD')
    
    if not sender_email or not sender_password:
        return {
            'success': False,
            'message': 'Email credentials not configured. Set SENDER_EMAIL and SENDER_PASSWORD in .env'
        }
    
    try:
        msg = MIMEMultipart('alternative')
        msg['Subject'] = f"Ledger Report — {datetime.now().strftime('%Y-%m-%d')}"
        msg['From'] = sender_email
        msg['To'] = recipient_email
        
        s = report['summary']
        html_body = f"""
        <html>
            <head></head>
            <body style="font-family: Arial, sans-serif; color: #333;">
                <h1 style="color: #1F3A93;">Ledger — AP Reconciliation Report</h1>
                <p>Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
                
                <h2 style="color: #1F3A93; margin-top: 30px;">Executive Summary</h2>
                <table style="border-collapse: collapse; width: 100%; max-width: 500px;">
                    <tr style="background: #1F3A93; color: white;">
                        <th style="padding: 10px; text-align: left;">Metric</th>
                        <th style="padding: 10px; text-align: left;">Value</th>
                    </tr>
                    <tr style="background: #F9FAFB;">
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">Total Lines Processed</td>
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">{s['total_lines_processed']}</td>
                    </tr>
                    <tr>
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">Total Exceptions</td>
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">{s['total_exceptions']}</td>
                    </tr>
                    <tr style="background: #F9FAFB;">
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">Amount at Risk</td>
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">${s['total_amount_at_risk']:,.2f}</td>
                    </tr>
                    <tr>
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">Touchless Rate</td>
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">{s['touchless_rate_pct']}%</td>
                    </tr>
                    <tr style="background: #F9FAFB;">
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">Needing Review</td>
                        <td style="padding: 8px; border: 1px solid #E5E7EB;">{s['needs_review']}</td>
                    </tr>
                </table>
                
                <p style="margin-top: 30px; color: #6B7280;">
                    For detailed exception analysis and policy citations, please log into the dashboard or review the attached Excel spreadsheet.
                </p>
            </body>
        </html>
        """
        msg.attach(MIMEText(html_body, 'html'))
        
        pdf_bytes = export_to_pdf(report)
        pdf_attachment = MIMEApplication(pdf_bytes)
        pdf_attachment.add_header('Content-Disposition', 'attachment', filename='report.pdf')
        msg.attach(pdf_attachment)
        
        excel_bytes = export_to_excel(report)
        excel_attachment = MIMEApplication(excel_bytes)
        excel_attachment.add_header('Content-Disposition', 'attachment', filename='report.xlsx')
        msg.attach(excel_attachment)
        
        with smtplib.SMTP(smtp_server, smtp_port) as server:
            server.starttls()
            server.login(sender_email, sender_password)
            server.send_message(msg)
        
        return {
            'success': True,
            'message': f'Report successfully sent to {recipient_email}'
        }
        
    except Exception as e:
        return {
            'success': False,
            'message': f'Error sending email: {str(e)}'
        }
